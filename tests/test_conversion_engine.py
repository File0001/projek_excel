"""Automated verification test for Excel to PDF converter core engine.

Tests metadata inspection, Excel COM session safety, and PDF generation with
various page setups (A4, F4, Portrait, Landscape, Fit scaling).
"""

import os
import sys

# Ensure project root is in sys.path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import openpyxl
from openpyxl.styles import Font, PatternFill

from src.config import PaperSizeEnum, PageOrientationEnum, PageScalingEnum
from src.core.excel_inspector import inspectWorkbook
from src.core.process_guard import ExcelSession
from src.core.excel_converter import convertExcelToPdf, ConversionOptions
from src.utils.logger import appLogger


def createSampleWorkbook(filePath: str) -> None:
    """Creates a sample multi-sheet Excel file with real data and styling for testing."""
    wb = openpyxl.Workbook()

    # Sheet 1: Sales Summary
    ws1 = wb.active
    ws1.title = "Sales Summary"
    ws1["A1"] = "Laporan Penjualan Bulanan"
    ws1["A1"].font = Font(bold=True, size=14)
    ws1.append(["Bulan", "Target", "Realisasi", "Persentase"])
    ws1.append(["Januari", 10000000, 12500000, "=C3/B3"])
    ws1.append(["Februari", 12000000, 11000000, "=C4/B4"])
    ws1.append(["Maret", 15000000, 16200000, "=C5/B5"])

    # Sheet 2: Department Expenses
    ws2 = wb.create_sheet("Department Expenses")
    ws2["A1"] = "Rincian Biaya Operasional"
    ws2["A1"].font = Font(bold=True, size=14)
    ws2.append(["Departemen", "Anggaran", "Terpakai", "Sisa"])
    ws2.append(["IT", 50000000, 42000000, "=B3-C3"])
    ws2.append(["Marketing", 30000000, 29500000, "=B4-C4"])
    ws2.append(["HR", 15000000, 13800000, "=B5-C5"])

    # Sheet 3: Hidden Reference
    ws3 = wb.create_sheet("Hidden Config")
    ws3["A1"] = "App Secret Config"
    ws3.sheet_state = "hidden"

    wb.save(filePath)
    wb.close()
    appLogger.info(f"Test workbook dibuat: {filePath}")


def testCoreEngine() -> None:
    """Executes end-to-end verification of the core conversion engine."""
    appLogger.info("=== MEMULAI TEST CORE ENGINE ===")

    testDir = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "test_output"))
    os.makedirs(testDir, exist_ok=True)

    testExcelPath = os.path.join(testDir, "test_input.xlsx")
    testSinglePdfPath = os.path.join(testDir, "test_single_sheet_a4.pdf")
    testMergedPdfPath = os.path.join(testDir, "test_merged_f4_landscape.pdf")

    try:
        # Step 1: Create sample workbook
        createSampleWorkbook(testExcelPath)
        assert os.path.exists(testExcelPath), "Gagal membuat sample Excel file."

        # Step 2: Test Inspection
        sheets = inspectWorkbook(testExcelPath)
        assert len(sheets) == 3, f"Expected 3 sheets, got {len(sheets)}"
        sheetNames = [s.name for s in sheets]
        assert "Sales Summary" in sheetNames
        assert "Department Expenses" in sheetNames
        assert "Hidden Config" in sheetNames
        appLogger.info("[PASS] Test Metadata Inspection Berhasil!")

        # Step 3: Test Single Sheet A4 Portrait Fit-Sheet
        singleOptions = ConversionOptions(
            inputExcelPath=testExcelPath,
            outputPdfPath=testSinglePdfPath,
            targetSheetNames=["Sales Summary"],
            paperSize=PaperSizeEnum.A4,
            orientation=PageOrientationEnum.PORTRAIT,
            scaling=PageScalingEnum.FIT_SHEET,
            mergeToSinglePdf=True,
        )
        outSingle = convertExcelToPdf(singleOptions)
        assert os.path.exists(outSingle), "PDF single sheet tidak terbentuk."
        assert os.path.getsize(outSingle) > 1000, "Ukuran file PDF terlalu kecil (indikasi rusak)."
        appLogger.info(f"[PASS] Test Single Sheet A4 PDF Berhasil! Size: {os.path.getsize(outSingle)} bytes")

        # Step 4: Test Multi-Sheet Merged F4 Landscape Fit-All-Columns
        mergedOptions = ConversionOptions(
            inputExcelPath=testExcelPath,
            outputPdfPath=testMergedPdfPath,
            targetSheetNames=["Sales Summary", "Department Expenses"],
            paperSize=PaperSizeEnum.F4,
            orientation=PageOrientationEnum.LANDSCAPE,
            scaling=PageScalingEnum.FIT_ALL_COLUMNS,
            mergeToSinglePdf=True,
        )
        outMerged = convertExcelToPdf(mergedOptions)
        assert os.path.exists(outMerged), "PDF merged multi-sheet tidak terbentuk."
        assert os.path.getsize(outMerged) > 1000, "Ukuran file PDF merged terlalu kecil."
        appLogger.info(f"[PASS] Test Merged Multi-Sheet F4 PDF Berhasil! Size: {os.path.getsize(outMerged)} bytes")

        appLogger.info("=== SELURUH TEST CORE ENGINE SUKSES TANPA ERROR ===")

    finally:
        # Cleanup test artifacts
        for tempFile in [testExcelPath, testSinglePdfPath, testMergedPdfPath]:
            if os.path.exists(tempFile):
                try:
                    os.remove(tempFile)
                except Exception:
                    pass
        if os.path.exists(testDir):
            try:
                os.rmdir(testDir)
            except Exception:
                pass


if __name__ == "__main__":
    testCoreEngine()
