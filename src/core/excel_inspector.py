"""Fast workbook metadata and sheet inspector.

Supports .xlsx and .xlsm via openpyxl (read-only mode),
with automatic fallback to COM inspector for legacy .xls files.
Adheres to camelCase variable naming and modular design.
"""

from dataclasses import dataclass
import os
from typing import List
import openpyxl

from src.config import supportedExtensions
from src.core.process_guard import ExcelSession
from src.utils.logger import appLogger


@dataclass
class SheetMetadata:
    """Metadata representing a single worksheet in a workbook."""
    name: str
    index: int
    isVisible: bool
    state: str


class EncryptedWorkbookError(Exception):
    """Raised when an Excel file is password-protected or encrypted."""
    pass


class InvalidWorkbookError(Exception):
    """Raised when an Excel file is corrupted or unsupported."""
    pass


def inspectWorkbookOpenpyxl(filePath: str) -> List[SheetMetadata]:
    """Inspects sheets using openpyxl for fast read-only metadata retrieval."""
    sheetList: List[SheetMetadata] = []
    workbook = None
    try:
        workbook = openpyxl.load_workbook(
            filename=filePath,
            read_only=True,
            data_only=True,
            keep_vba=True,
        )
        for sheetIndex, sheetName in enumerate(workbook.sheetnames):
            worksheet = workbook[sheetName]
            # openpyxl sheet_state can be 'visible', 'hidden', or 'veryHidden'
            sheetState = getattr(worksheet, "sheet_state", "visible")
            isVisible = (sheetState == "visible")
            sheetList.append(
                SheetMetadata(
                    name=sheetName,
                    index=sheetIndex + 1,
                    isVisible=isVisible,
                    state=sheetState,
                )
            )
        return sheetList
    except openpyxl.utils.exceptions.InvalidFileException as invalidErr:
        raise InvalidWorkbookError(f"Format file tidak valid: {invalidErr}")
    except Exception as inspectErr:
        errText = str(inspectErr).lower()
        if "encrypted" in errText or "password" in errText:
            raise EncryptedWorkbookError("File Excel terenkripsi atau dilindungi kata sandi.")
        raise
    finally:
        if workbook is not None:
            try:
                workbook.close()
            except Exception:
                pass


def inspectWorkbookCom(filePath: str) -> List[SheetMetadata]:
    """Inspects sheets using Excel COM instance (used for legacy .xls or complex files)."""
    sheetList: List[SheetMetadata] = []
    normalizedPath = os.path.abspath(filePath)

    with ExcelSession() as excelApp:
        workbookCom = None
        try:
            workbookCom = excelApp.Workbooks.Open(
                normalizedPath,
                ReadOnly=True,
                UpdateLinks=0,
                IgnoreReadOnlyRecommended=True,
            )
            sheetsCount = workbookCom.Worksheets.Count
            for sheetIndex in range(1, sheetsCount + 1):
                sheetItem = workbookCom.Worksheets(sheetIndex)
                # xlSheetVisible = -1, xlSheetHidden = 0, xlSheetVeryHidden = 2
                isVisible = (sheetItem.Visible == -1)
                stateStr = "visible" if isVisible else "hidden"
                sheetList.append(
                    SheetMetadata(
                        name=sheetItem.Name,
                        index=sheetIndex,
                        isVisible=isVisible,
                        state=stateStr,
                    )
                )
            return sheetList
        except Exception as comErr:
            errStr = str(comErr).lower()
            if "password" in errStr or "encrypted" in errStr:
                raise EncryptedWorkbookError("File Excel terenkripsi atau memerlukan kata sandi.")
            raise InvalidWorkbookError(f"Gagal membaca sheet workbook melalui COM: {comErr}")
        finally:
            if workbookCom is not None:
                try:
                    workbookCom.Close(SaveChanges=False)
                except Exception:
                    pass


def inspectWorkbook(filePath: str) -> List[SheetMetadata]:
    """Main entrypoint for workbook inspection with intelligent engine selection.

    Validates file existence and extension, attempts ultra-fast openpyxl inspection,
    and falls back to COM automation for legacy .xls formats.
    """
    if not os.path.exists(filePath):
        raise FileNotFoundError(f"File tidak ditemukan: {filePath}")

    _, fileExtension = os.path.splitext(filePath)
    normalizedExt = fileExtension.lower()

    if normalizedExt not in supportedExtensions:
        raise ValueError(
            f"Ekstensi file '{fileExtension}' tidak didukung. Didukung: {', '.join(supportedExtensions)}"
        )

    appLogger.info(f"Menginspeksi metadata workbook: {filePath}")

    # Legacy .xls binary format requires COM inspection
    if normalizedExt == ".xls":
        appLogger.info("Format .xls terdeteksi, menggunakan COM inspector.")
        return inspectWorkbookCom(filePath)

    # Modern .xlsx and .xlsm formats
    try:
        return inspectWorkbookOpenpyxl(filePath)
    except EncryptedWorkbookError:
        raise
    except Exception as fallbackErr:
        appLogger.warning(
            f"openpyxl gagal ({fallbackErr}), mencoba inspeksi via Excel COM..."
        )
        return inspectWorkbookCom(filePath)
